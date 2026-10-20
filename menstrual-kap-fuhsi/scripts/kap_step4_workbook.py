"""FUHSI KAP study — Phase 4: Excel workbook (13 sheets, 10 native charts).
Run:  .venv/bin/python kap_step4_workbook.py
Input: data/processed/kap_cleaned_full.csv, reports/output/tables/*.csv
Output: reports/output/KAP_Menstrual_Irregularities_FUHSI_Analysis.xlsx
"""
import warnings
from pathlib import Path
from collections import Counter
import numpy as np
import pandas as pd
import openpyxl as xl

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import BarChart, PieChart, DoughnutChart, Reference

warnings.filterwarnings("ignore")

BASE = Path(__file__).resolve().parent
TAB = BASE / "reports" / "output" / "tables"
CLEAN = BASE / "data" / "processed" / "kap_cleaned_full.csv"
OUT_XLSX = BASE / "reports" / "output" / "KAP_Menstrual_Irregularities_FUHSI_Analysis.xlsx"

TITLE_FONT = Font(bold=True, size=14, color="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF")
SUB_FONT = Font(bold=True, size=11, color="1F4E79")
NOTE_FONT = Font(italic=True, size=9, color="555555")
HDR_FILL = PatternFill("solid", fgColor="1F4E79")
ALT_FILL = PatternFill("solid", fgColor="DDEBF7")

def rd(name):
    d = pd.read_csv(TAB / name, keep_default_na=False)
    for c in d.columns:
        if d[c].dtype == object and len(d) and d[c].str.fullmatch(r"-?\d+(\.\d+)?").all():
            d[c] = pd.to_numeric(d[c])
    return d

df = pd.read_csv(CLEAN, dtype=str, keep_default_na=False)
num_cols = ["age_num", "level_num", "know_score", "know_pct", "att_score", "type_n"] \
    + [f"A{i+1}__n" for i in range(10)] + [f"E{i+1}__n" for i in range(14)] \
    + [f"hbm_{g}" for g in ["susceptibility", "severity", "benefit", "barrier", "cues", "self_efficacy"]]
for c in num_cols:
    df[c] = pd.to_numeric(df[c], errors="coerce")
N = len(df)

def sheet_df(wb, name, d, widths=None, note=None, freeze="A2"):
    ws = wb.create_sheet(name)
    for r, row in enumerate(dataframe_to_rows(d, index=False), 1):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    for c in range(1, d.shape[1] + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = HDR_FONT
        cell.fill = HDR_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for w, col in enumerate(d.columns, 1):
        letter = get_column_letter(w)
        if widths and col in widths:
            ws.column_dimensions[letter].width = widths[col]
        else:
            ws.column_dimensions[letter].width = min(max(12, int(d[col].astype(str).str.len().max() if len(d) else 12) + 2), 60)
    ws.freeze_panes = freeze
    if note:
        ws.cell(row=d.shape[0] + 3, column=1, value=note).font = NOTE_FONT
    return ws

wb = Workbook()
wb.remove(wb.active)

# ============================ README ============================
ws = wb.create_sheet("README")
readme = [
    ("FUHSI KAP Study — Knowledge, Attitude and Health-Seeking Behaviour Regarding Menstrual Irregularities", "title"),
    ("Client: Taoreed Sofiat Opemiposi (22/NSC/048), Department of Nursing Science", "bold"),
    (f"Sample: {N} consenting female undergraduate respondents (377 administered, 1 non-consent excluded)", "bold"),
    ("", ""),
    ("ANALYSIS PER THESIS DOCUMENT (Ch.1 objectives + §3.11 analysis plan)", "sub"),
    ("README → Data_Dictionary → Cleaned_Data (376 records + scores) → Demographics → Obj1_Knowledge → Obj2_Attitude → Obj3_HSB → Obj4_HBM → Chi_Square → Pearson → Supplementary → Charts_Native (10 native Excel charts) → Key_Findings.", ""),
    ("", ""),
    ("Primary analyses (document-mandated)", "sub"),
    ("Descriptive: frequencies/percentages, mean/SD for all items and scores.", ""),
    ("Knowledge & attitude categorised by mean AND median split (both coincide at these data: cut = 11 on knowledge 0-12, 41 on attitude 10-50).", ""),
    ("Chi-square tests: H1 knowledge x future help-seeking intent; H2 attitude x future intent (+ secondary HSB associations with ever-sought-help and demographics, Cramer's V).", ""),
    ("Pearson correlation: knowledge, attitude, HBM constructs, age, level.", ""),
    ("", ""),
    ("Supplementary (clearly labelled, NOT in §3.11)", "sub"),
    ("S01 Bloom grades (>70% good) + >=50% adequacy; S02 Spearman matrix; S03 Kruskal-Wallis/Mann-Whitney group comparisons; S04 logistic regression on future intent.", ""),
    ("", ""),
    ("DISCLOSED DEVIATIONS (documented in report §limitations)", "sub"),
    ("1. Likert responses collected on a 5-point scale (SA/A/U/D/SD) although the proposal §3.7 describes a 4-point scale — reported as fielded.", ""),
    ("2. Reliability: pilot test-retest (§3.9) was outside the analyst's remit; Cronbach's alpha computed on the main-study data instead (knowledge KR-20 = 0.77, attitude alpha = 0.43 [low — treat attitudes dimensionally], HBM total = 0.82).", ""),
    ("3. Item wording as fielded: Section A includes ethnicity (no religion item); Q2 'type of irregularity' was administered in two form editions (both response columns merged).", ""),
    ("4. Form branching was not enforced in the field: 28 respondents who reported no irregularity also listed types, and 52 non-experienced respondents answered Q3/Q4 — experiential tables are restricted to the experienced subset.", ""),
]
r = 1
for txt, kind in readme:
    cell = ws.cell(row=r, column=1, value=txt)
    if kind == "title": cell.font = TITLE_FONT
    elif kind == "bold": cell.font = SUB_FONT
    elif kind == "sub": cell.font = PatternFill() or SUB_FONT  # placeholder, replaced below
    if kind == "sub":
        cell.font = SUB_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
ws.column_dimensions["A"].width = 150

# ============================ Data_Dictionary ============================
rows = [
    ("Timestamp", "Survey submission timestamp (raw)"), ("consent_raw?", "raw responses kept left of scored columns"),
]
DICT = [
    ("Timestamp", "raw", "Submission timestamp"),
    ("age_num", "derived", "Exact age in completed years (parsed from free text; 164 exact, 212 banded only)"),
    ("age_band", "derived", "Age band (16-20 / 21-24 / 25-30 / Above 30 / Not stated)"),
    ("age_grp", "derived", "Age group for chi-square (16-20 / 21-24 / 25+)"),
    ("level_c / level_num", "raw+derived", "Level 100-400 (categorical / numeric)"),
    ("dept_c / dept_group", "derived", "Department (normalised) / grouped Clinical (MBBS, Nursing, MLS) vs Non-clinical"),
    ("marital_c / ethnicity_c / menarche_c", "raw+clean", "Marital status, ethnicity, age at menarche band"),
    ("K1..K12 (+_raw, __n)", "Section B", "12 knowledge True/False/Don't-know items; __n = 1 correct / 0 incorrect (key: all True except K10=False)"),
    ("know_score / know_pct", "derived", "Knowledge score 0-12 and percentage"),
    ("know_cat_mean / know_cat_median", "§3.11", "Good vs Poor by mean (>=10.63) and median (>=11) split"),
    ("A1..A10 (+_raw, __n)", "Section C", "10 attitude items 1-5 (5=SA..1=SD); items 1,2,4,10 reverse-scored"),
    ("att_score", "derived", "Attitude score 10-50"),
    ("att_cat_mean / att_cat_median", "§3.11", "Positive vs Negative by mean (>=40.39) / median (>=41) split"),
    ("E1..E14 (+_raw, __n)", "Section E", "14 HBM items 1-5"),
    ("hbm_susceptibility / severity / benefit / barrier / cues / self_efficacy", "derived", "HBM construct means (barrier: higher = more perceived barriers)"),
    ("exp_c", "Q1", "Ever experienced menstrual irregularity (Yes/No)"),
    ("type_list / type_n", "Q2", "Types reported (multi-select merged across both form editions; ; separated)"),
    ("action_c", "Q3", "First action after noticing irregularity (experienced subset)"),
    ("time_c / sought_ever", "Q4", "Timeliness of seeking professional help / ever sought professional help"),
    ("info_list", "Q5", "Information sources (multi-select)"),
    ("campus_c", "Q6", "Ever visited FUHSI campus health centre"),
    ("facility_c", "Q7", "Preferred facility for menstrual care"),
    ("barrier_list", "Q8", "Barriers to seeking care (multi-select)"),
    ("future_c / future2", "Q9", "Would seek professional help in future (3-cat / binary Yes vs Not sure/No)"),
    ("k/u raw name columns", "raw", "Columns ending _raw preserve the verbatim response of the matching item"),
]
sheet_df(wb, "Data_Dictionary", pd.DataFrame(DICT, columns=["column(s)", "section", "description"]),
         widths={"column(s)": 40, "section": 12, "description": 110})

# ============================ Cleaned_Data ============================
sheet_df(wb, "Cleaned_Data", df, widths={"Timestamp": 24}, note="376 consenting respondents; multi-select fields ';'-separated lists.")

# ============================ Demographics ============================
t01 = rd("T01a_demographics_categorical.csv")
t01b = rd("T01b_departments.csv")
ws = wb.create_sheet("Demographics")
ws.cell(row=1, column=1, value="TABLE 1 — Demographic profile (n=376)").font = TITLE_FONT
for r, row in enumerate(dataframe_to_rows(t01, index=False), 3):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
r2 = t01.shape[0] + 5
ws.cell(row=r2, column=1, value="Department detail").font = SUB_FONT
for r, row in enumerate(dataframe_to_rows(t01b, index=False), r2 + 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
for c in range(1, 6):
    ws.column_dimensions[get_column_letter(c)].width = 34

# ============================ Obj1_Knowledge ============================
ws = wb.create_sheet("Obj1_Knowledge")
ws.cell(row=1, column=1, value="OBJECTIVE 1 — Level of knowledge").font = TITLE_FONT
t02a, t02b, t02c = rd("T02a_knowledge_items.csv"), rd("T02b_knowledge_scores.csv"), rd("T02c_knowledge_summary.csv")
ws.cell(row=2, column=1, value="Item-level results").font = SUB_FONT
for r, row in enumerate(dataframe_to_rows(t02a, index=False), 3):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
r2 = t02a.shape[0] + 5
ws.cell(row=r2, column=1, value="Score summary & categorisation (mean and median splits)").font = SUB_FONT
for r, row in enumerate(dataframe_to_rows(pd.concat([t02c, t02b], ignore_index=True).fillna(""), index=False), r2 + 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
for c, w in [(1, 10), (2, 90), (3, 14), (4, 12), (5, 12), (6, 12), (7, 12), (8, 14), (9, 14)]:
    ws.column_dimensions[get_column_letter(c)].width = w

# ============================ Obj2_Attitude ============================
ws = wb.create_sheet("Obj2_Attitude")
ws.cell(row=1, column=1, value="OBJECTIVE 2 — Attitudes (5-point Likert as fielded)").font = TITLE_FONT
t03a, t03b = rd("T03a_attitude_items.csv"), rd("T03b_attitude_scores.csv")
for r, row in enumerate(dataframe_to_rows(t03a, index=False), 2):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
r2 = t03a.shape[0] + 4
ws.cell(row=r2, column=1, value="Attitude score summary & categorisation").font = SUB_FONT
for r, row in enumerate(dataframe_to_rows(t03b, index=False), r2 + 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
ws.cell(row=r2 + t03b.shape[0] + 2, column=1,
        value="Note: items A1, A2, A4, A10 are negatively worded and reverse-scored before summing (10-50).").font = NOTE_FONT
for c, w in [(1, 8), (2, 90), (3, 12), (4, 10), (5, 10), (6, 10), (7, 10), (8, 10), (9, 10), (10, 10), (11, 10), (12, 10)]:
    ws.column_dimensions[get_column_letter(c)].width = w

# ============================ Obj3_HSB ============================
ws = wb.create_sheet("Obj3_HSB")
ws.cell(row=1, column=1, value="OBJECTIVE 3 — Health-seeking behaviour patterns").font = TITLE_FONT
r = 2
blocks = [("T05a_experienced.csv", "Q1 Ever experienced menstrual irregularity"),
          ("T05b_types.csv", "Q2 Types experienced (multi-select; base = 285 respondents listing types)"),
          ("T05c_first_action.csv", "Q3 First action (experienced subset n=257)"),
          ("T05d_timeliness.csv", "Q4 Timeliness of professional help-seeking (experienced subset)"),
          ("T05e_info_sources.csv", "Q5 Information sources (multi-select, % of all 376)"),
          ("T05f_campus_health_center.csv", "Q6 Visited FUHSI campus health centre"),
          ("T05g_preferred_facility.csv", "Q7 Preferred facility for menstrual care"),
          ("T05h_future_intent.csv", "Q9 Would seek professional help in future"),
          ("T05i_barriers.csv", "Q8 Main barriers to seeking care (multi-select, % of all 376)")]
for f, ttl in blocks:
    t = rd(f)
    ws.cell(row=r, column=1, value=ttl).font = SUB_FONT
    for rr, row in enumerate(dataframe_to_rows(t, index=False), r + 1):
        for c, v in enumerate(row, 1):
            ws.cell(row=rr, column=c, value=v)
    r += t.shape[0] + 3
for c in range(1, 6):
    ws.column_dimensions[get_column_letter(c)].width = 52

# ============================ Obj4_HBM ============================
ws = wb.create_sheet("Obj4_HBM")
ws.cell(row=1, column=1, value="OBJECTIVE 4 / Section E — HBM constructs and perceived barriers").font = TITLE_FONT
t04a, t04b = rd("T04a_hbm_items.csv"), rd("T04b_hbm_subscales.csv")
for r, row in enumerate(dataframe_to_rows(t04b, index=False), 2):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
r2 = t04b.shape[0] + 4
ws.cell(row=r2, column=1, value="Item-level results").font = SUB_FONT
for r, row in enumerate(dataframe_to_rows(t04a, index=False), r2 + 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
ws.cell(row=r2 + t04a.shape[0] + 2, column=1,
        value="Note: barrier items E7-E10 scored so higher = greater perceived barrier; cues alpha low (0.32) — 2-item scale.").font = NOTE_FONT
for c, w in [(1, 8), (2, 90), (3, 18), (4, 10), (5, 10), (6, 10), (7, 10), (8, 10), (9, 10), (10, 10), (11, 10), (12, 10)]:
    ws.column_dimensions[get_column_letter(c)].width = w

# ============================ Chi_Square ============================
ws = wb.create_sheet("Chi_Square")
ws.cell(row=1, column=1, value="TABLE — Chi-square tests (χ², df, p, Cramér's V)").font = TITLE_FONT
t06a, t06b = rd("T06a_hypotheses_chi_square.csv"), rd("T06b_demographic_chi_square.csv")
ws.cell(row=2, column=1, value="Hypothesis tests (document-mandated)").font = SUB_FONT
for r, row in enumerate(dataframe_to_rows(t06a, index=False), 3):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
r2 = t06a.shape[0] + 5
ws.cell(row=r2, column=1, value="Associations with demographics").font = SUB_FONT
for r, row in enumerate(dataframe_to_rows(t06b, index=False), r2 + 1):
    for c, v in enumerate(row, 1):
        ws.cell(row=r, column=c, value=v)
for c in range(1, 12):
    ws.column_dimensions[get_column_letter(c)].width = 24
ws.column_dimensions["A"].width = 30

# ============================ Pearson ============================
ws = wb.create_sheet("Pearson")
ws.cell(row=1, column=1, value="TABLE — Pearson correlations (pairwise)").font = TITLE_FONT
r = 2
for f, ttl in [("T07a_pearson_r.csv", "r coefficients"),
               ("T07b_pearson_p.csv", "p-values"),
               ("T07c_pearson_n.csv", "pairwise n")]:
    t = rd(f)
    ws.cell(row=r, column=1, value=ttl).font = SUB_FONT
    for rr, row in enumerate(dataframe_to_rows(t, index=False), r + 1):
        for c, v in enumerate(row, 1):
            ws.cell(row=rr, column=c, value=v)
    r += t.shape[0] + 3
for c in range(1, 12):
    ws.column_dimensions[get_column_letter(c)].width = 16

# ============================ Supplementary ============================
ws = wb.create_sheet("Supplementary")
ws.cell(row=1, column=1, value="SUPPLEMENTARY ANALYSES — beyond §3.11, clearly labelled").font = TITLE_FONT
r = 2
for f, ttl, w_a in [("S01_bloom_grades_supplementary.csv", "S01 Bloom grading of knowledge", 40),
                    ("S03_group_comparisons_supplementary.csv", "S03 Kruskal-Wallis / Mann-Whitney (group comparisons)", 40),
                    ("S04_logistic_future_intent_supplementary.csv", "S04 Logistic regression — future intent (OR, 95% CI)", 40)]:
    t = rd(f)
    ws.cell(row=r, column=1, value=ttl).font = SUB_FONT
    for rr, row in enumerate(dataframe_to_rows(t, index=False), r + 1):
        for c, v in enumerate(row, 1):
            ws.cell(row=rr, column=c, value=v)
    r += t.shape[0] + 3
for c in range(1, 8):
    ws.column_dimensions[get_column_letter(c)].width = 30

# ============================ Charts_Native (10 native charts) ============================
ws = wb.create_sheet("Charts_Native")
ws.column_dimensions["A"].width = 34
ws.column_dimensions["B"].width = 16
ws.column_dimensions["C"].width = 4

def add_chart_block(row, title, cats, vals, chart):
    ws.cell(row=row, column=1, value=title).font = SUB_FONT
    ws.cell(row=row + 1, column=1, value="Category").font = HDR_FONT
    ws.cell(row=row + 1, column=2, value="Value").font = HDR_FONT
    ws.cell(row=row + 1, column=1).fill = HDR_FILL
    ws.cell(row=row + 1, column=2).fill = HDR_FILL
    for i, (c, v) in enumerate(zip(cats, vals)):
        ws.cell(row=row + 2 + i, column=1, value=str(c))
        ws.cell(row=row + 2 + i, column=2, value=round(float(v), 1))
    n = len(cats)
    data = Reference(ws, min_col=2, min_row=row + 1, max_row=row + 1 + n)
    labs = Reference(ws, min_col=1, min_row=row + 2, max_row=row + 1 + n)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(labs)
    chart.height = 7.5
    chart.width = 14
    ws.add_chart(chart, f"D{row}")
    return max(n + 4, 22)

row = 1
c1 = DoughnutChart(); c1.title = "Ever experienced menstrual irregularity"; c1.style = 10
eu = df["exp_c"].value_counts()
row += add_chart_block(row, "C1 — Ever experienced (counts)", eu.index.tolist(), eu.values.tolist(), c1)

c2 = PieChart(); c2.title = "Would seek help in future"; c2.style = 10
fc = df["future_c"].value_counts()
row += add_chart_block(row, "C2 — Future intent (counts)", fc.index.tolist(), fc.values.tolist(), c2)

c3 = PieChart(); c3.title = "Knowledge category (mean split)"; c3.style = 10
kc = df["know_cat_mean"].value_counts()
row += add_chart_block(row, "C3 — Knowledge categorisation", kc.index.tolist(), kc.values.tolist(), c3)

c4 = DoughnutChart(); c4.title = "Attitude category (mean split)"; c4.style = 10
ac = df["att_cat_mean"].value_counts()
row += add_chart_block(row, "C4 — Attitude categorisation", ac.index.tolist(), ac.values.tolist(), c4)

c5 = BarChart(); c5.type = "bar"; c5.title = "Knowledge items % correct"; c5.style = 10
t02a = rd("T02a_knowledge_items.csv")
row += add_chart_block(row, "C5 — Knowledge items (% correct)", t02a["item"].tolist(),
                       t02a["correct_pct"].tolist(), c5)

c6 = BarChart(); c6.type = "bar"; c6.title = "Attitude item means (1-5, reversed recoded)"; c6.style = 10
t03a = rd("T03a_attitude_items.csv")
row += add_chart_block(row, "C6 — Attitude item means", t03a["item"].tolist(),
                       t03a["mean_1to5"].tolist(), c6)

c7 = BarChart(); c7.type = "col"; c7.title = "HBM construct means (1-5)"; c7.style = 10
t04b = rd("T04b_hbm_subscales.csv")
row += add_chart_block(row, "C7 — HBM constructs", t04b["subscale"].tolist(),
                       t04b["mean"].tolist(), c7)

c8 = BarChart(); c8.type = "bar"; c8.title = "Types of irregularity (mentions)"; c8.style = 10
t05b = rd("T05b_types.csv")
row += add_chart_block(row, "C8 — Types reported (%)", t05b["type_of_irregularity"].tolist(),
                       t05b["pct_of_285_reporting"].tolist(), c8)

c9 = BarChart(); c9.type = "bar"; c9.title = "Information sources (mentions %)"; c9.style = 10
t05e = rd("T05e_info_sources.csv")
row += add_chart_block(row, "C9 — Info sources (% of 376)", t05e["source"].tolist(),
                       t05e["pct_of_376"].tolist(), c9)

c10 = BarChart(); c10.type = "bar"; c10.title = "Barriers to seeking care (mentions %)"; c10.style = 10
t05i = rd("T05i_barriers.csv")
row += add_chart_block(row, "C10 — Barriers (% of 376)", t05i["barrier"].tolist(),
                       t05i["pct_of_376"].tolist(), c10)

# ============================ Key_Findings ============================
ws = wb.create_sheet("Key_Findings")
t06a = rd("T06a_hypotheses_chi_square.csv")
h1b = t06a[t06a["test"] == "H1 primary (binary)"].iloc[0]
h2b = t06a[t06a["test"] == "H2 primary (binary)"].iloc[0]
t05a = rd("T05a_experienced.csv")
t05b = rd("T05b_types.csv")
t05c = rd("T05c_first_action.csv")
t05d = rd("T05d_timeliness.csv")
t05e = rd("T05e_info_sources.csv")
t05f = rd("T05f_campus_health_center.csv")
t05g = rd("T05g_preferred_facility.csv")
t05h = rd("T05h_future_intent.csv")
t05i = rd("T05i_barriers.csv")
t04b = rd("T04b_hbm_subscales.csv")
t01 = rd("T01a_demographics_categorical.csv")
s04 = rd("S04_logistic_future_intent_supplementary.csv")

exp_yes_n = int(t05a.loc[t05a["category"] == "Yes", "n"].iloc[0])
fut_yes_n = int(t05h.loc[t05h["category"] == "Yes", "n"].iloc[0])
fut_yes_pct = float(t05h.loc[t05h["category"] == "Yes", "pct"].iloc[0])
never = int(t05d.loc[(t05d["variable"] == "time_to_seek_help") & (t05d["value"] == "Did not seek"), "n"].iloc[0])
sought = exp_yes_n - never
campus_yes = int(t05f.loc[t05f["category"] == "Yes", "n"].iloc[0])
top_action = None
for c in ["Ignored it / did nothing"]:
    m = t05c[t05c["category"] == c]
    if len(m): top_action = (m.iloc[0]["n"], m.iloc[0]["pct"])
type_top = t05b.iloc[0]
bar_top = t05i.iloc[0]
info_top = t05e.iloc[0]
know_good = int((df["know_cat_mean"] == "Good").sum())
att_pos = int((df["att_cat_mean"] == "Positive").sum())

kf = [
    ("KEY FINDINGS", "title"),
    (f"1. Knowledge: mean {df['know_score'].mean():.2f}/12 (median 11); Good knowledge {know_good} ({know_good/N*100:.1f}%) vs Poor {N-know_good} by both mean and median split. Weakest items: K10 (normalisation myth), K1 (cycle length), sect K3/K8 (see Obj1 sheet)."),
    (f"2. Attitudes: mean {df['att_score'].mean():.2f}/50 (median 41) — Positive {att_pos} ({att_pos/N*100:.1f}%). Stigma/embarrassment salient (A2 mean below midpoint; alpha 0.43 → attitude treated dimensionally)."),
    (f"3. Experience: {exp_yes_n} ({exp_yes_n/N*100:.1f}%) ever experienced menstrual irregularity; top types: {type_top['type_of_irregularity']} ({type_top['mentions']} mentions, {type_top['pct_of_285_reporting']}% of 285 reporting), then dysmenorrhea, amenorrhea, menorrhagia."),
    (f"4. First response (experienced n={exp_yes_n}): largest group took no action ({top_action[0]}, {top_action[1]}%); only sought professional care: {sought} ({sought/exp_yes_n*100:.1f}%), and {never} ({never/exp_yes_n*100:.1f}%) never sought help."),
    (f"5. Campus health centre: only {campus_yes} ({campus_yes/N*100:.1f}%) ever visited it; preferred facility overall is private clinic ({int(t05g.loc[t05g['category']=='Private clinic','n'].iloc[0])})."),
    (f"6. Future intent: {fut_yes_n} ({fut_yes_pct}%) would seek professional help in future.", ""),
    (f"7. Information: {info_top['source']} most used ({info_top['pct_of_376']}% of 376)."),
    (f"8. Barriers: '{bar_top['barrier']}' dominant ({bar_top['pct_of_376']}% of 376); HBM barrier construct mean {t04b.loc[t04b['subscale']=='barrier','mean'].iloc[0]}/5."),
    (f"9. H1 (knowledge x future intent): χ²(1)={h1b['chi2']}, p={h1b['p']}, V={h1b['cramers_v']} — significant, weak: Good knowledge more likely to intend seeking (91% vs 81%). H0a rejected."),
    (f"10. H2 (attitude x future intent): χ²(1)={h2b['chi2']}, p={h2b['p']}, V={h2b['cramers_v']} — not significant; H0b retained."),
    ("11. Demographics (chi-square): knowledge by level (p=0.001, V=0.21) and clinical vs non-clinical dept (p<0.001, V=0.33); future intent by level (p=0.002). Attitude had no significant demographic associations."),
    ("12. Pearson: knowledge correlates with self-efficacy (r=0.50), severity (r=0.44), susceptibility (r=0.32); attitude negatively with barriers (r=-0.26); age-knowledge r=0.28; level negatively with HBM constructs (junior students more risk-averse perceptions) — all p<0.05 candidates; see Pearson sheet."),
    ("13. Supplementary: logistic (n=376) confirms knowledge (OR 1.22/p=0.005] and level (OR 0.59 per 100-level step/p=0.002) predict future intent.", ""),
]
r = 1
for item in kf:
    txt = item[0]
    kind = item[1] if len(item) > 1 else ""
    cell = ws.cell(row=r, column=1, value=txt)
    cell.font = TITLE_FONT if kind == "title" else Font(bold=kind == "bold")
    if kind == "sub": cell.font = SUB_FONT
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r].height = 30 if kind != "title" else 22
    r += 1
ws.column_dimensions["A"].width = 140

# tab colors
for nm, col in [("README", "1F4E79"), ("Data_Dictionary", "808080"), ("Cleaned_Data", "808080"),
                ("Demographics", "2E75B6"), ("Obj1_Knowledge", "2E75B6"), ("Obj2_Attitude", "2E75B6"),
                ("Obj3_HSB", "2E75B6"), ("Obj4_HBM", "2E75B6"), ("Chi_Square", "C00000"),
                ("Pearson", "C00000"), ("Supplementary", "808080"), ("Charts_Native", "548235"),
                ("Key_Findings", "548235")]:
    if nm in wb.sheetnames:
        wb[nm].sheet_properties.tabColor = col

order = ["README", "Data_Dictionary", "Cleaned_Data", "Demographics", "Obj1_Knowledge",
         "Obj2_Attitude", "Obj3_HSB", "Obj4_HBM", "Chi_Square", "Pearson", "Supplementary",
         "Charts_Native", "Key_Findings"]
wb._sheets = [wb[n] for n in order]
wb.save(OUT_XLSX)
print("Saved workbook:", OUT_XLSX.name, "| sheets:", wb.sheetnames, "| charts: 10")
